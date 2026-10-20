from __future__ import annotations

import csv
import io
from pathlib import Path

import pymupdf
from docx import Document as WordDocument
from openpyxl import load_workbook

from . import ai
from .config import PAGES


def _record(kind: str, locator: str, text: str, **metadata) -> dict:
    return {"kind": kind, "locator": locator, "text": text.strip(), "metadata": metadata}


def parse_pdf(path: Path, document_id: str) -> tuple[list[dict], int, str | None]:
    elements: list[dict] = []
    unreadable = 0
    pdf = pymupdf.open(path)
    page_dir = PAGES / document_id
    page_dir.mkdir(parents=True, exist_ok=True)
    try:
        if pdf.needs_pass:
            raise ValueError("PDF protégé par mot de passe")
        if len(pdf) > 300:
            raise ValueError("PDF de plus de 300 pages : divisez-le avant l'import")
        for number, page in enumerate(pdf, 1):
            blocks = page.get_text("blocks")
            page_text = " ".join(str(block[4]) for block in blocks if len(block) > 4)
            for index, block in enumerate(blocks):
                if len(block) < 5 or not str(block[4]).strip():
                    continue
                text = str(block[4]).strip()
                elements.append(_record("paragraph", f"p.{number} bloc {index + 1}", text,
                                        page=number, bbox=[round(float(v), 1) for v in block[:4]], parser="native_pdf"))
            if len(page_text.strip()) < 80:
                image = page_dir / f"page-{number}.png"
                page.get_pixmap(matrix=pymupdf.Matrix(1.5, 1.5), alpha=False).save(image)
                if ai.enabled("vision"):
                    try:
                        visual_text = ai.describe_page(image)
                    except Exception as exc:
                        elements.append(_record("warning", f"p.{number}", f"Lecture visuelle échouée : {exc}", page=number))
                        unreadable += 1
                    else:
                        if visual_text.strip():
                            elements.append(_record("visual", f"p.{number}", visual_text, page=number, parser="vlm"))
                        else:
                            unreadable += 1
                else:
                    unreadable += 1
        note = f"{unreadable} page(s) sans texte exploitable ; configurez un VLM pour les lire." if unreadable else None
        return elements, len(pdf), note
    finally:
        pdf.close()


def parse_docx(path: Path) -> tuple[list[dict], int, str | None]:
    doc = WordDocument(path)
    elements: list[dict] = []
    for i, para in enumerate(doc.paragraphs, 1):
        if para.text.strip():
            elements.append(_record("paragraph", f"paragraphe {i}", para.text,
                                    style=para.style.name if para.style else ""))
    for table_no, table in enumerate(doc.tables, 1):
        for row_no, row in enumerate(table.rows, 1):
            text = " | ".join(cell.text.replace("\n", " / ") for cell in row.cells)
            if text.strip(" |"):
                elements.append(_record("table_row", f"tableau {table_no} ligne {row_no}", text))
    return elements, 1, None


def parse_xlsx(path: Path) -> tuple[list[dict], int, str | None]:
    formulas = load_workbook(path, read_only=True, data_only=False)
    values = load_workbook(path, read_only=True, data_only=True)
    elements: list[dict] = []
    try:
        for sheet in formulas:
            if sheet.max_row > 20000 or sheet.max_column > 200:
                raise ValueError("Classeur trop grand pour cette version (20 000 lignes, 200 colonnes par onglet)")
            value_sheet = values[sheet.title]
            for row in sheet.iter_rows():
                for cell in row:
                    if cell.value is None:
                        continue
                    value = value_sheet[cell.coordinate].value
                    text = str(value if value is not None else cell.value)
                    if isinstance(cell.value, str) and cell.value.startswith("="):
                        text += f" [formule : {cell.value}]"
                    elements.append(_record("cell", f"{sheet.title}!{cell.coordinate}", text,
                                            sheet=sheet.title, cell=cell.coordinate, formula=str(cell.value) if cell.data_type == "f" else None))
        return elements, len(formulas.sheetnames), None
    finally:
        formulas.close()
        values.close()


def parse_csv(path: Path) -> tuple[list[dict], int, str | None]:
    raw = path.read_bytes()
    text = raw.decode("utf-8-sig", errors="replace")
    try:
        dialect = csv.Sniffer().sniff(text[:4096])
    except csv.Error:
        dialect = csv.excel
    elements = []
    for row_no, row in enumerate(csv.reader(io.StringIO(text), dialect), 1):
        if row_no > 20000:
            raise ValueError("CSV de plus de 20 000 lignes : divisez-le avant l'import")
        if any(value.strip() for value in row):
            elements.append(_record("table_row", f"ligne {row_no}", " | ".join(row)))
    return elements, 1, None


def parse_image(path: Path, document_id: str) -> tuple[list[dict], int, str | None]:
    if not ai.enabled("vision"):
        return [], 1, "Image non lue : configurez AI_API_KEY et AI_VISION_MODEL."
    text = ai.describe_page(path)
    return ([_record("visual", "image entière", text, parser="vlm")] if text.strip() else []), 1, None


def parse(path: Path, ext: str, document_id: str) -> tuple[list[dict], int, str | None]:
    if ext == ".pdf":
        return parse_pdf(path, document_id)
    if ext == ".docx":
        return parse_docx(path)
    if ext == ".xlsx":
        return parse_xlsx(path)
    if ext == ".csv":
        return parse_csv(path)
    if ext in {".png", ".jpg", ".jpeg"}:
        return parse_image(path, document_id)
    raise ValueError(f"Format non pris en charge : {ext}")

