"""Create synthetic documents for a local demonstration. No client data."""
from pathlib import Path

import pymupdf
from docx import Document
from openpyxl import Workbook

ROOT = Path(__file__).resolve().parent.parent
DEST = ROOT / "demo-documents"
DEST.mkdir(exist_ok=True)

pdf = pymupdf.open()
page = pdf.new_page()
page.insert_textbox(
    pymupdf.Rect(50, 60, 550, 600),
    "SOCIETE DEMO - Politique de sauvegarde\n\n"
    "Une sauvegarde quotidienne des bases de donnees de production est prevue. "
    "Le responsable infrastructure effectue une revue mensuelle des journaux de sauvegarde.\n\n"
    "Ce document est une politique et ne constitue pas une preuve de test de restauration.",
    fontsize=12,
)
pdf.save(DEST / "politique_sauvegarde.pdf")
pdf.close()

doc = Document()
doc.add_heading("Société Démo - procédure de gestion des incidents", 0)
doc.add_paragraph("Les incidents de sécurité sont enregistrés dans un registre. Une procédure d'escalade vers l'équipe sécurité est définie.")
doc.add_paragraph("Le document ne contient pas de preuve que cette procédure a été testée.")
doc.save(DEST / "procedure_incidents.docx")

book = Workbook()
sheet = book.active
sheet.title = "Revue des accès"
sheet.append(["Contrôle", "Fréquence", "Responsable", "Résultat"])
sheet.append(["Revue des droits administrateur", "trimestrielle", "RSSI", "À réaliser"])
sheet.append(["Inventaire des actifs", "annuelle", "DSI", "Documenté"])
book.save(DEST / "registre_controles.xlsx")

print(f"Documents fictifs créés dans {DEST}")

