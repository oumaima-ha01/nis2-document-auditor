import io
import json
import time
from pathlib import Path

import pymupdf
from docx import Document
from fastapi.testclient import TestClient
from openpyxl import Workbook

from nis2_auditor import ai, catalog, config, db, engine, main, parsers
from nis2_auditor.catalog import applicability, load_catalog


def test_applicability_requires_confirmed_entity_type():
    req = load_catalog()["requirements"][0]
    assert applicability(req, {"entity_type": "unknown"})[0] == "undetermined"
    assert applicability(req, {"entity_type": "important"})[0] == "applicable"


def test_ai_citation_is_rejected_when_not_in_source():
    candidate = {"id": "item", "text": "Une politique de sauvegarde existe."}
    raw = [{"element_id": "item", "quote": "Des restaurations sont testées chaque mois.", "explanation": "x"}]
    assert engine._valid_evidence(raw, {"item": candidate}) == []


def test_ai_api_contract_without_external_request(tmp_path, monkeypatch):
    monkeypatch.setattr(config, "AI_API_KEY", "test-key")
    monkeypatch.setattr(config, "AI_TEXT_MODEL", "text-test")
    monkeypatch.setattr(config, "AI_VISION_MODEL", "vision-test")
    calls = []

    class FakeResponse:
        def raise_for_status(self):
            pass

        def json(self):
            content = {"text": "Sauvegarde quotidienne"} if calls[-1]["model"] == "vision-test" else {
                "status": "uncertain", "rationale": "Politique seule", "evidence": []}
            return {"choices": [{"message": {"content": json.dumps(content)}}]}

    def fake_post(url, headers, json, timeout):
        assert url.endswith("/chat/completions")
        assert headers["Authorization"] == "Bearer test-key"
        assert json["response_format"] == {"type": "json_object"}
        calls.append(json)
        return FakeResponse()

    monkeypatch.setattr(ai.httpx, "post", fake_post)
    result = ai.assess(load_catalog()["requirements"][0], [{"id": "one", "locator": "p.1", "text": "PSSI"}])
    assert result["status"] == "uncertain"
    image_path = tmp_path / "image.png"
    image_path.write_bytes(b"\x89PNG\r\n\x1a\n")
    assert ai.describe_page(image_path) == "Sauvegarde quotidienne"
    assert calls[1]["messages"][1]["content"][1]["image_url"]["url"].startswith("data:image/png;base64,")


def test_recyf_catalog_is_complete_and_readable(monkeypatch):
    monkeypatch.setattr(catalog, "RECYF_PATH", Path("missing-recyf-catalog.json"))
    monkeypatch.setenv("CATALOG_MODE", "recyf")
    requirements = load_catalog()["requirements"]
    assert len(requirements) == 152
    assert len({row["id"] for row in requirements}) == 152
    assert all(row["summary"] and "\ufffd" not in row["summary"] for row in requirements)
    assert any(row["essential_only"] for row in requirements)


def test_document_parsers_cover_word_excel_csv_and_scanned_pdf(tmp_path, monkeypatch):
    monkeypatch.setattr(parsers, "PAGES", tmp_path / "pages")
    monkeypatch.setattr(config, "AI_API_KEY", "")

    word_path = tmp_path / "procedure.docx"
    word = Document()
    word.add_paragraph("Procédure de notification des incidents")
    table = word.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Responsable"
    table.cell(0, 1).text = "RSSI"
    word.save(word_path)
    word_elements, _, _ = parsers.parse(word_path, ".docx", "word")
    assert any("notification" in item["text"] for item in word_elements)
    assert any(item["kind"] == "table_row" for item in word_elements)

    excel_path = tmp_path / "registre.xlsx"
    workbook = Workbook()
    workbook.active["A1"] = "Sauvegarde"
    workbook.save(excel_path)
    excel_elements, _, _ = parsers.parse(excel_path, ".xlsx", "excel")
    assert any(item["locator"] == "Sheet!A1" and item["text"] == "Sauvegarde" for item in excel_elements)

    csv_path = tmp_path / "registre.csv"
    csv_path.write_text("Contrôle;État\nMFA;Actif\n", encoding="utf-8")
    csv_elements, _, _ = parsers.parse(csv_path, ".csv", "csv")
    assert any("MFA" in item["text"] for item in csv_elements)

    pdf_path = tmp_path / "scan.pdf"
    pdf = pymupdf.open()
    pdf.new_page()
    pdf.save(pdf_path)
    pdf.close()
    pdf_elements, _, note = parsers.parse(pdf_path, ".pdf", "scan")
    assert not pdf_elements
    assert "1 page(s)" in note
    assert (tmp_path / "pages" / "scan" / "page-1.png").exists()


def test_full_local_flow(tmp_path, monkeypatch):
    files = tmp_path / "files"
    pages = tmp_path / "pages"
    files.mkdir()
    pages.mkdir()
    monkeypatch.setattr(db, "DB", tmp_path / "test.sqlite3")
    monkeypatch.setattr(main, "FILES", files)
    monkeypatch.setattr(engine, "FILES", files)
    monkeypatch.setattr(parsers, "PAGES", pages)
    monkeypatch.setattr(config, "AI_API_KEY", "")
    monkeypatch.setattr(catalog, "RECYF_PATH", Path("missing-recyf-catalog.json"))
    with TestClient(main.app) as client:
        audit = client.post("/api/audits", json={"name": "Test local", "entity_type": "important"}).json()
        pdf = pymupdf.open()
        page = pdf.new_page()
        page.insert_text((40, 70), "Politique de sauvegarde et continuité des systèmes d'information")
        page.insert_text((40, 100), "La sauvegarde quotidienne est prévue dans la politique de sécurité.")
        stream = io.BytesIO(pdf.tobytes())
        pdf.close()
        uploaded = client.post(f"/api/audits/{audit['id']}/documents", files={"file": ("politique.pdf", stream, "application/pdf")})
        assert uploaded.status_code == 200
        for _ in range(80):
            docs = client.get(f"/api/audits/{audit['id']}/documents").json()
            if docs[0]["status"] not in {"queued", "processing"}:
                break
            time.sleep(.05)
        assert docs[0]["status"] == "ready", docs
        assert client.post(f"/api/audits/{audit['id']}/run").status_code == 200
        for _ in range(80):
            status = client.get(f"/api/audits/{audit['id']}").json()["status"]
            if status != "processing":
                break
            time.sleep(.05)
        assert status == "ready"
        matrix = client.get(f"/api/audits/{audit['id']}/matrix").json()
        backup = next(x for x in matrix if x["id"] == "NIS2-21-2-C")
        assert backup["assessment"]["status"] == "uncertain"
        assert backup["evidence"]
        evidence = backup["evidence"][0]
        assert evidence["locator"].startswith("p.1")
        review = client.post(f"/api/assessments/{backup['assessment']['id']}/review",
                             json={"status": "partially_supported", "note": "Politique trouvée, test non prouvé."})
        assert review.status_code == 200
        assert client.get(f"/api/audits/{audit['id']}/export.csv").status_code == 200
        exported = client.get(f"/api/audits/{audit['id']}/export.json")
        assert exported.status_code == 200
        assert exported.json()["audit"]["id"] == audit["id"]
        assert client.get(f"/api/elements/{evidence['element_id']}/preview.png").headers["content-type"] == "image/png"
        monkeypatch.setenv("CATALOG_MODE", "recyf")
        assert client.post(f"/api/audits/{audit['id']}/run").status_code == 200
        for _ in range(100):
            status = client.get(f"/api/audits/{audit['id']}").json()["status"]
            if status != "processing":
                break
            time.sleep(.05)
        assert status == "ready"
        full_matrix = client.get(f"/api/audits/{audit['id']}/matrix").json()
        assert len(full_matrix) == 152
        assert any(row["assessment"]["status"] == "not_applicable" for row in full_matrix)
        extra = client.post(f"/api/audits/{audit['id']}/documents",
                            files={"file": ("registre.csv", b"controle;etat\nMFA;actif\n", "text/csv")})
        assert extra.status_code == 200
        assert all(row["assessment"] is None for row in client.get(f"/api/audits/{audit['id']}/matrix").json())

